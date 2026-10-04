# sync_index.py — recorre K: (SOLO LECTURA), extrae texto y actualiza
# SQLite FTS5 + manifest.json de forma incremental. No llama a ningún LLM.
import hashlib
import io
import json
import os
import re
import sys
import time
from datetime import datetime

import config
from dtr_common import abrir_db, normalizar, slugify, extracto_de, pid_vivo

PREFIJO_LARGO = "\\\\?\\"          # soporte de rutas >260 chars en Windows


def ruta_larga(p: str) -> str:
    return p if p.startswith(PREFIJO_LARGO) else PREFIJO_LARGO + p


def fingerprint(path: str, size: int) -> str:
    """Hash rápido: primeros y últimos 64 KB + tamaño."""
    h = hashlib.blake2b(digest_size=16)
    h.update(str(size).encode())
    with open(ruta_larga(path), "rb") as f:
        h.update(f.read(65536))
        if size > 131072:
            f.seek(size - 65536)
            h.update(f.read(65536))
    return h.hexdigest()


# ---------------- extractores ----------------

def _limpia(t: str) -> str:
    t = re.sub(r"[ \t]+", " ", t or "")
    t = re.sub(r"\n{3,}", "\n\n", t)
    t = t.strip()[:config.MAX_TEXTO_CHARS]
    # PDFs corruptos pueden producir surrogates que SQLite/UTF-8 rechazan
    return t.encode("utf-8", errors="replace").decode("utf-8")


def extraer_pdf(path: str, size: int):
    if size > config.MAX_PDF_BYTES:
        return "", None, "grande-omitido"
    from pypdf import PdfReader
    with open(ruta_larga(path), "rb") as f:
        reader = PdfReader(f, strict=False)
        npag = len(reader.pages)
        partes = []
        total = 0
        for i, page in enumerate(reader.pages):
            if i >= config.MAX_PDF_PAGINAS or total > config.MAX_TEXTO_CHARS:
                break
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            partes.append(t)
            total += len(t)
    texto = _limpia("\n".join(partes))
    if len(texto) < config.MIN_TEXTO_PDF:
        return texto, npag, "escaneado(ocr?)"
    return texto, npag, "ok"


def extraer_docx(path: str, size: int):
    import docx
    d = docx.Document(ruta_larga(path))
    partes = [p.text for p in d.paragraphs if p.text]
    for tabla in d.tables[:50]:
        for fila in tabla.rows[:200]:
            partes.append(" | ".join(c.text for c in fila.cells))
    return _limpia("\n".join(partes)), None, "ok"


def extraer_xlsx(path: str, size: int):
    import openpyxl
    wb = openpyxl.load_workbook(ruta_larga(path), read_only=True,
                                data_only=True)
    partes = []
    try:
        for hoja in wb.worksheets[:10]:
            partes.append(f"[Hoja: {hoja.title}]")
            for i, fila in enumerate(hoja.iter_rows(values_only=True)):
                if i >= config.MAX_XLSX_FILAS:
                    break
                celdas = [str(c) for c in fila if c is not None]
                if celdas:
                    partes.append(" | ".join(celdas))
    finally:
        wb.close()
    return _limpia("\n".join(partes)), None, "ok"


def extraer_plano(path: str, size: int):
    with open(ruta_larga(path), "rb") as f:
        raw = f.read(config.MAX_TEXTO_CHARS)
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return _limpia(raw.decode(enc)), None, "ok"
        except UnicodeDecodeError:
            continue
    return _limpia(raw.decode("utf-8", errors="replace")), None, "ok"


def _html_a_texto(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    html = re.sub(r"(?s)<[^>]+>", " ", html)
    import html as h
    return h.unescape(html)


def extraer_html(path: str, size: int):
    texto, _, _ = extraer_plano(path, size)
    return _limpia(_html_a_texto(texto)), None, "ok"


def extraer_eml(path: str, size: int):
    import email
    from email import policy
    with open(ruta_larga(path), "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    cab = [f"Asunto: {msg.get('subject', '')}", f"De: {msg.get('from', '')}",
           f"Para: {msg.get('to', '')}", f"Fecha: {msg.get('date', '')}"]
    cuerpo = ""
    try:
        parte = msg.get_body(preferencelist=("plain", "html"))
        if parte:
            cuerpo = parte.get_content()
            if parte.get_content_type() == "text/html":
                cuerpo = _html_a_texto(cuerpo)
    except Exception:
        pass
    adjuntos = [fn for fn in
                (p.get_filename() for p in msg.walk()) if fn]
    if adjuntos:
        cab.append("Adjuntos: " + ", ".join(adjuntos[:20]))
    return _limpia("\n".join(cab) + "\n\n" + (cuerpo or "")), None, "ok"


def extraer_msg(path: str, size: int):
    import extract_msg
    m = extract_msg.openMsg(ruta_larga(path))
    try:
        cab = [f"Asunto: {m.subject or ''}", f"De: {m.sender or ''}",
               f"Para: {m.to or ''}", f"Fecha: {m.date or ''}"]
        cuerpo = m.body or ""
        adj = [a.longFilename or a.shortFilename or "" for a in m.attachments]
        if adj:
            cab.append("Adjuntos: " + ", ".join(a for a in adj[:20] if a))
    finally:
        m.close()
    return _limpia("\n".join(cab) + "\n\n" + cuerpo), None, "ok"


EXTRACTORES = {
    ".pdf": extraer_pdf, ".docx": extraer_docx, ".xlsx": extraer_xlsx,
    ".txt": extraer_plano, ".md": extraer_plano, ".csv": extraer_plano,
    ".html": extraer_html, ".htm": extraer_html,
    ".eml": extraer_eml, ".msg": extraer_msg,
}


def procesar_archivo(path_abs: str, ext: str, size: int):
    """Devuelve (texto, paginas, estado). Nunca lanza excepción."""
    try:
        if ext in EXTRACTORES:
            texto, npag, estado = EXTRACTORES[ext](path_abs, size)
            if estado == "ok" and not texto:
                estado = "sin-texto"
            return texto, npag, estado
        return "", None, "sin-extractor"
    except Exception as e:
        return "", None, f"error:{type(e).__name__}"


# ---------------- recorrido incremental ----------------

def cargar_manifest(con=None):
    try:
        with open(config.MANIFEST_PATH, encoding="utf-8") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        pass
    # recuperación tras crash: reconstruir desde la base de datos
    if con is not None:
        m = {r["ruta_rel"]: {"hash": r["hash"], "mtime": r["mtime"],
                             "size": r["size"]}
             for r in con.execute("SELECT ruta_rel, hash, mtime, size "
                                  "FROM docs")}
        if m:
            print(f"manifest.json ausente: reconstruido desde la BD "
                  f"({len(m)} entradas)", flush=True)
        return m
    return {}


def guardar_manifest(manifest):
    os.makedirs(config.META_DIR, exist_ok=True)
    tmp = config.MANIFEST_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False)
    os.replace(tmp, config.MANIFEST_PATH)


def recorrer_raiz():
    """Genera (ruta_abs_sin_prefijo, ruta_rel, expediente, ext, size, mtime)."""
    raiz = ruta_larga(config.RAIZ_DATOS.rstrip("\\") + "\\")
    for dirpath, dirnames, filenames in os.walk(raiz, onerror=lambda e: None):
        dirnames[:] = [d for d in dirnames
                       if normalizar(d) not in config.DIRS_EXCLUIDOS]
        for fn in filenames:
            ext = os.path.splitext(fn)[1].lower()
            if ext not in config.EXT_TEXTO and ext not in config.EXT_FICHA:
                continue
            abs_pref = os.path.join(dirpath, fn)
            try:
                st = os.stat(abs_pref)
            except OSError:
                continue
            abs_limpia = abs_pref[len(PREFIJO_LARGO):] \
                if abs_pref.startswith(PREFIJO_LARGO) else abs_pref
            rel = os.path.relpath(abs_limpia, config.RAIZ_DATOS)
            partes = rel.split(os.sep)
            expediente = partes[0] if len(partes) > 1 else ""
            yield abs_limpia, rel, expediente, ext, st.st_size, st.st_mtime


def indexar_uno(cur, manifest, slug_de, rel, expediente, abs_p, ext, size,
                mtime, h, stats, slugs_dirty):
    prev = manifest.get(rel)
    texto, npag, estado = procesar_archivo(abs_p, ext, size)
    if estado.startswith("error"):
        stats["errores"] += 1
    slug = slug_de(expediente)
    nombre = os.path.basename(rel)
    ahora = datetime.now().isoformat(timespec="seconds")
    cur.execute("""
        INSERT INTO docs(ruta_rel, expediente, slug, nombre, nombre_norm,
                         ext, size, mtime, hash, paginas, estado,
                         extracto, actualizado)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(ruta_rel) DO UPDATE SET
          expediente=excluded.expediente, slug=excluded.slug,
          nombre=excluded.nombre, nombre_norm=excluded.nombre_norm,
          ext=excluded.ext, size=excluded.size, mtime=excluded.mtime,
          hash=excluded.hash, paginas=excluded.paginas,
          estado=excluded.estado, extracto=excluded.extracto,
          actualizado=excluded.actualizado
    """, (rel, expediente or "(raíz)", slug, nombre, normalizar(nombre),
          ext, size, mtime, h, npag, estado, extracto_de(texto), ahora))
    doc_id = cur.execute("SELECT id FROM docs WHERE ruta_rel=?",
                         (rel,)).fetchone()["id"]
    cur.execute("DELETE FROM textos WHERE doc_id=?", (doc_id,))
    if texto:
        cur.execute("INSERT INTO textos(doc_id, texto) VALUES (?,?)",
                    (doc_id, texto))
    stats["modificados" if prev else "nuevos"] += 1
    slugs_dirty.add(slug)
    manifest[rel] = {"hash": h, "mtime": mtime, "size": size}


def sync():
    t0 = time.time()
    con = abrir_db()
    cur = con.cursor()
    manifest = cargar_manifest(con)
    vistos = set()
    stats = {"nuevos": 0, "modificados": 0, "sin_cambios": 0,
             "borrados": 0, "errores": 0}
    slugs_dirty = set()

    # mapa expediente -> slug estable (colisiones -> sufijo)
    nombres_exp = {}

    def slug_de(expediente: str) -> str:
        if expediente == "":
            return "_raiz"
        if expediente not in nombres_exp:
            base = slugify(expediente)
            usados = set(nombres_exp.values())
            slug, n = base, 2
            while slug in usados:
                slug, n = f"{base}-{n}", n + 1
            nombres_exp[expediente] = slug
        return nombres_exp[expediente]

    # precarga slugs existentes para estabilidad entre corridas
    for row in cur.execute("SELECT slug, nombre FROM expedientes"):
        nombres_exp[row["nombre"]] = row["slug"]

    procesados = 0
    for abs_p, rel, expediente, ext, size, mtime in recorrer_raiz():
        vistos.add(rel)
        prev = manifest.get(rel)
        if prev and prev["size"] == size and abs(prev["mtime"] - mtime) < 2:
            stats["sin_cambios"] += 1
            continue
        try:
            h = fingerprint(abs_p, size)
        except OSError:
            stats["errores"] += 1
            continue
        if prev and prev.get("hash") == h:
            manifest[rel] = {"hash": h, "mtime": mtime, "size": size}
            stats["sin_cambios"] += 1
            continue

        try:
            indexar_uno(cur, manifest, slug_de, rel, expediente, abs_p,
                        ext, size, mtime, h, stats, slugs_dirty)
        except Exception as e:
            print(f"  !! error irrecuperable en {rel!r}: "
                  f"{type(e).__name__}", flush=True)
            stats["errores"] += 1

        procesados += 1
        if procesados % 2000 == 0:
            con.commit()
            guardar_manifest(manifest)
            print(f"  ... {procesados} procesados "
                  f"({time.time()-t0:,.0f}s)", flush=True)
    # borrados: en manifest pero ya no en K:
    for rel in sorted(set(manifest) - vistos):
        row = cur.execute("SELECT id, slug FROM docs WHERE ruta_rel=?",
                          (rel,)).fetchone()
        if row:
            cur.execute("DELETE FROM textos WHERE doc_id=?", (row["id"],))
            cur.execute("DELETE FROM docs WHERE id=?", (row["id"],))
            slugs_dirty.add(row["slug"])
        del manifest[rel]
        stats["borrados"] += 1

    # tabla de expedientes + marcado dirty
    cur.execute("""
        CREATE TEMP TABLE agg AS
        SELECT slug, expediente AS nombre, COUNT(*) AS n, SUM(size) AS s
        FROM docs GROUP BY slug
    """)
    ahora = datetime.now().isoformat(timespec="seconds")
    for row in cur.execute("SELECT * FROM agg").fetchall():
        ruta_abs = config.RAIZ_DATOS if row["slug"] == "_raiz" else \
            os.path.join(config.RAIZ_DATOS, row["nombre"])
        cur.execute("""
            INSERT INTO expedientes(slug, nombre, nombre_norm, ruta_abs,
                                    n_docs, size_total, actualizado, dirty)
            VALUES (?,?,?,?,?,?,?,1)
            ON CONFLICT(slug) DO UPDATE SET
              n_docs=excluded.n_docs, size_total=excluded.size_total,
              actualizado=CASE WHEN expedientes.n_docs<>excluded.n_docs
                   OR expedientes.size_total<>excluded.size_total
                   THEN excluded.actualizado ELSE expedientes.actualizado END
        """, (row["slug"], row["nombre"], normalizar(row["nombre"]),
              ruta_abs, row["n"], row["s"], ahora))
    # expedientes desaparecidos
    cur.execute("DELETE FROM expedientes WHERE slug NOT IN "
                "(SELECT slug FROM agg)")
    if slugs_dirty:
        cur.executemany("UPDATE expedientes SET dirty=1 WHERE slug=?",
                        [(s,) for s in slugs_dirty])
    con.commit()

    grafo_stats = reconstruir_grafo(con)

    cur.execute("INSERT INTO meta(k, v) VALUES ('ultima_sync', ?) "
                "ON CONFLICT(k) DO UPDATE SET v=excluded.v",
                (datetime.now().isoformat(timespec="seconds"),))
    con.commit()

    guardar_manifest(manifest)

    dur = time.time() - t0
    print(f"SYNC OK en {dur/60:,.1f} min — nuevos={stats['nuevos']} "
          f"modificados={stats['modificados']} borrados={stats['borrados']} "
          f"sin_cambios={stats['sin_cambios']} errores={stats['errores']} "
          f"aristas={grafo_stats}", flush=True)
    if not any((stats["nuevos"], stats["modificados"], stats["borrados"])):
        print("Sin cambios en K: — cerebro ya al día.", flush=True)
    con.close()
    return stats


# ---------------- grafo mecánico ----------------

def _consulta_nombre(nombre: str):
    """Deriva una consulta FTS a partir del nombre de un expediente, o None
    si el nombre es demasiado genérico para buscarlo sin ruido."""
    tokens = [t for t in re.findall(r"[a-z0-9]+", normalizar(nombre))
              if t not in config.STOPWORDS_GRAFO and not t.isdigit()
              and len(t) >= 3]
    tokens = tokens[:4]
    if len(tokens) >= 2:
        cuerpo = " ".join(f'"{t}"' for t in tokens)
        return f"NEAR({cuerpo}, 12)"
    if len(tokens) == 1 and len(tokens[0]) >= 6:
        return f'"{tokens[0]}"'
    return None


def reconstruir_grafo(con) -> int:
    cur = con.cursor()
    antes = {(r["src"], r["dst"], r["tipo"])
             for r in cur.execute("SELECT src, dst, tipo FROM aristas")}
    cur.execute("DELETE FROM aristas")

    expedientes = cur.execute(
        "SELECT slug, nombre FROM expedientes WHERE slug != '_raiz'"
    ).fetchall()

    # 1) menciona-a: docs de otros expedientes que citan el nombre del cliente
    for exp in expedientes:
        q = _consulta_nombre(exp["nombre"])
        if not q:
            continue
        try:
            hits = cur.execute("""
                SELECT d.slug AS src, COUNT(*) AS peso,
                       MIN(d.nombre) AS ejemplo
                FROM fts JOIN docs d ON d.id = fts.rowid
                WHERE fts MATCH ? AND d.slug != ?
                GROUP BY d.slug ORDER BY peso DESC LIMIT ?
            """, (q, exp["slug"], config.MAX_ARISTAS_POR_EXP)).fetchall()
        except Exception:
            continue
        cur.executemany("""
            INSERT INTO aristas(src, dst, tipo, peso, ejemplo)
            VALUES (?,?,?,?,?)
            ON CONFLICT(src, dst, tipo) DO UPDATE SET
              peso=excluded.peso, ejemplo=excluded.ejemplo
        """, [(h["src"], exp["slug"], "menciona-a", h["peso"], h["ejemplo"])
              for h in hits])

    # 2) nombre-similar: expedientes que comparten un token distintivo
    por_token = {}
    for exp in expedientes:
        for t in set(re.findall(r"[a-z]+", normalizar(exp["nombre"]))):
            if len(t) >= 6 and t not in config.STOPWORDS_GRAFO:
                por_token.setdefault(t, []).append(exp["slug"])
    for token, slugs in por_token.items():
        if 2 <= len(slugs) <= 8:
            for a in slugs:
                for b in slugs:
                    if a < b:
                        cur.execute("""
                            INSERT INTO aristas(src, dst, tipo, peso, ejemplo)
                            VALUES (?,?,?,1,?)
                            ON CONFLICT(src, dst, tipo) DO UPDATE SET
                              peso=aristas.peso+1,
                              ejemplo=aristas.ejemplo || ', ' || excluded.ejemplo
                        """, (a, b, "nombre-similar", token))

    despues = {(r["src"], r["dst"], r["tipo"])
               for r in cur.execute("SELECT src, dst, tipo FROM aristas")}
    tocados = {s for (s, d, t) in antes ^ despues} | \
              {d for (s, d, t) in antes ^ despues}
    if tocados:
        cur.executemany("UPDATE expedientes SET dirty=1 WHERE slug=?",
                        [(s,) for s in tocados])
    con.commit()
    return len(despues)


def adquirir_lock() -> bool:
    """Evita dos sincronizaciones simultáneas. True si obtuvimos el lock."""
    os.makedirs(config.META_DIR, exist_ok=True)
    lock = os.path.join(config.META_DIR, "sync.pid")
    try:
        with open(lock) as f:
            pid = int(f.read().strip() or 0)
        if pid and pid_vivo(pid):
            print(f"Otra sincronización en curso (PID {pid}); salgo.",
                  flush=True)
            return False
    except (FileNotFoundError, ValueError):
        pass
    with open(lock, "w") as f:
        f.write(str(os.getpid()))
    return True


def liberar_lock():
    try:
        os.remove(os.path.join(config.META_DIR, "sync.pid"))
    except OSError:
        pass


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    print(f"=== sync_index {datetime.now():%Y-%m-%d %H:%M:%S} "
          f"raiz={config.RAIZ_DATOS}", flush=True)
    if not adquirir_lock():
        sys.exit(9)
    try:
        sync()
    finally:
        liberar_lock()
